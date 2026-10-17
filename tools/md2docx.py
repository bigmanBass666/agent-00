# -*- coding: utf-8 -*-
"""把项目内的 markdown 文档转成 DOCX（中文排版）"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement


def set_style_cjk_font(style, cn_font, en_font):
    style.font.name = en_font
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cn_font)
    rFonts.set(qn('w:ascii'), en_font)
    rFonts.set(qn('w:hAnsi'), en_font)


def clean_heading_style(style):
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
            if rFonts.get(qn(attr)):
                del rFonts.attrib[qn(attr)]
    pPr = style.element.find(qn('w:pPr'))
    if pPr is not None:
        pBdr = pPr.find(qn('w:pBdr'))
        if pBdr is not None:
            pPr.remove(pBdr)


def add_runs(para, text):
    """处理 **粗体** 与 `code`"""
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for p in parts:
        if not p:
            continue
        if p.startswith('**') and p.endswith('**'):
            r = para.add_run(p[2:-2]); r.bold = True
        elif p.startswith('`') and p.endswith('`'):
            r = para.add_run(p[1:-1]); r.font.name = 'Consolas'
        else:
            para.add_run(p)


def md_to_docx(md_path, docx_path, doc_title=None):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.8)

    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    set_style_cjk_font(normal, '宋体', 'Times New Roman')
    normal.paragraph_format.line_spacing = 1.4

    ts = doc.styles['Title']
    set_style_cjk_font(ts, '黑体', 'Arial')
    clean_heading_style(ts)
    ts.font.size = Pt(20)
    for i in range(1, 4):
        h = doc.styles[f'Heading {i}']
        set_style_cjk_font(h, '黑体', 'Arial')
        clean_heading_style(h)
        h.font.size = Pt({1: 15, 2: 13, 3: 11.5}[i])

    # 页脚页码
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for xml in [
        f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>',
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>',
        f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>',
    ]:
        r = fp.add_run(); r._r.append(__import__('docx').oxml.parse_xml(xml))

    i = 0
    in_code = False
    code_buf = []
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith('```'):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.line_spacing = 1.1
                r = p.add_run('\n'.join(code_buf))
                r.font.name = 'Consolas'
                r.font.size = Pt(9)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 表格
        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip('|').split('|')]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            t = doc.add_table(rows=len(rows) + 1, cols=len(header))
            t.style = 'Table Grid'
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ''
                p = cell.paragraphs[0]
                r = p.add_run(re.sub(r'\*\*', '', h)); r.bold = True; r.font.size = Pt(9.5)
                shd = __import__('docx').oxml.parse_xml(
                    f'<w:shd {nsdecls("w")} w:fill="E7E6E6"/>')
                cell._tc.get_or_add_tcPr().append(shd)
            for ri, row in enumerate(rows, start=1):
                for j, v in enumerate(row):
                    if j >= len(header):
                        break
                    cell = t.rows[ri].cells[j]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    add_runs(p, v)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = re.sub(r'\*\*', '', m.group(2))
            if level == 1 and first_h1:
                h = doc.add_heading(text, level=0)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_h1 = False
            else:
                doc.add_heading(text, level=min(level, 3))
            i += 1
            continue

        # 分隔线
        if stripped in ('---', '***'):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = __import__('docx').oxml.parse_xml(
                f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:color="BFBFBF"/></w:pBdr>')
            pPr.append(pBdr)
            i += 1
            continue

        # 引用
        if stripped.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_runs(p, stripped.lstrip('> ').strip())
            for r in p.runs:
                r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            i += 1
            continue

        # 列表
        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, m.group(1))
            i += 1
            continue
        m = re.match(r'^\d+[.)]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(1))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print('saved:', docx_path)


if __name__ == '__main__':
    md_to_docx(sys.argv[1], sys.argv[2])
